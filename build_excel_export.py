"""
Generate an Excel export of employee attrition predictions.

Run from the project root:
    python build_excel_export.py
"""

import pathlib

import joblib
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PROJECT_ROOT = pathlib.Path(__file__).resolve().parent
MODEL_PATH = PROJECT_ROOT / "models" / "best_attrition_model.pkl"
DATA_PATH = PROJECT_ROOT / "data" / "WA_Fn-UseC_-HR-Employee-Attrition.csv"
EXPORT_PATH = PROJECT_ROOT / "reports" / "Employee_Attrition_Predictions.xlsx"


def build_model_results(pipe, df: pd.DataFrame):
    df = add_age_group(df)
    drop = [
        "EmployeeCount",
        "EmployeeNumber",
        "Over18",
        "StandardHours",
        "Attrition",
    ]
    X = df.drop(columns=[c for c in drop if c in df.columns])
    X = X[pipe.feature_names_in_]

    y_true = df["Attrition"].map({"No": 0, "Yes": 1}).astype(int)
    y_pred = pipe.predict(X)
    y_prob = pipe.predict_proba(X)[:, 1]

    df_out = df.copy()
    df_out["Predicted_Attrition"] = np.where(y_pred == 1, "Yes", "No")
    df_out["Predicted_Prob_Leave"] = np.round(y_prob, 4)
    df_out["Prediction_Match"] = (df_out["Predicted_Attrition"] == df["Attrition"])

    return df_out, y_true, y_pred, y_prob


def make_summary_sheet(wb: Workbook, df: pd.DataFrame, y_true, y_pred, y_prob):
    from openpyxl.utils import get_column_letter

    summary_ws = wb.create_sheet("Prediction Summary")

    header_fill = PatternFill("solid", fgColor="1F3B63")
    header_font = Font(bold=True, color="FFFFFF")

    summary_ws["A1"] = "Metric"
    summary_ws["B1"] = "Value"
    summary_ws["A1"].fill = header_fill
    summary_ws["A1"].font = header_font
    summary_ws["B1"].fill = header_fill
    summary_ws["B1"].font = header_font

    rows = [
        ("Total Employees", int(len(df))),
        ("Actual Leavers (Attrition = Yes)", int((df["Attrition"] == "Yes").sum())),
        ("Predicted Leavers", int((y_pred == 1).sum())),
        ("Prediction Match Rate", float((y_pred == y_true).mean())),
        ("Mean Predicted Probability of Leaving", float(np.mean(y_prob))),
        ("Max Predicted Probability of Leaving", float(np.max(y_prob))),
        ("Min Predicted Probability of Leaving", float(np.min(y_prob))),
        ("Mean Predicted Prob for Actual Leavers", float(np.mean(y_prob[y_true == 1]))),
        ("Mean Predicted Prob for Actual Non-Leavers", float(np.mean(y_prob[y_true == 0]))),
    ]

    for i, (metric, value) in enumerate(rows, start=2):
        summary_ws[f"A{i}"] = metric
        summary_ws[f"B{i}"] = value

    summary_ws.column_dimensions["A"].width = 46
    summary_ws.column_dimensions["B"].width = 20

    return summary_ws


def load_dataset(path: pathlib.Path = DATA_PATH) -> pd.DataFrame:
    return pd.read_csv(path)


def add_age_group(df: pd.DataFrame) -> pd.DataFrame:
    # Match the AgeGroup binning used during training in the notebook.
    df = df.copy()
    df["AgeGroup"] = pd.cut(
        df["Age"],
        bins=[0, 25, 35, 45, 55, 100],
        labels=["<25", "25-35", "35-45", "45-55", "55+"],
        right=False,
    )
    return df


def predict_for_dataset(pipe, df: pd.DataFrame) -> pd.DataFrame:
    df = add_age_group(df)
    drop = [
        "EmployeeCount",
        "EmployeeNumber",
        "Over18",
        "StandardHours",
        "Attrition",
    ]
    X = df.drop(columns=[c for c in drop if c in df.columns])
    X = X[pipe.feature_names_in_]

    y_pred = pipe.predict(X)
    y_prob = pipe.predict_proba(X)[:, 1]

    out = df.copy()
    out["Predicted_Attrition"] = np.where(y_pred == 1, "Yes", "No")
    out["Predicted_Prob_Leave"] = np.round(y_prob, 4)
    out["Prediction_Match"] = (out["Predicted_Attrition"] == df["Attrition"])
    return out


def style_workbook(wb: Workbook, df: pd.DataFrame):
    header_fill = PatternFill("solid", fgColor="1F3B63")
    header_font = Font(bold=True, color="FFFFFF")
    yes_fill = PatternFill("solid", fgColor="F6C9C9")
    no_fill = PatternFill("solid", fgColor="D6E8D6")
    center = Alignment(horizontal="center", vertical="center")

    sheet = wb.active
    sheet.title = "Predictions"

    cols = list(df.columns)
    for cidx, col in enumerate(cols, start=1):
        cell = sheet.cell(row=1, column=cidx, value=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    for ridx, (_, row) in enumerate(df.iterrows(), start=2):
        for cidx, col in enumerate(cols, start=1):
            value = row[col]
            if pd.isna(value):
                value = ""
            cell = sheet.cell(row=ridx, column=cidx, value=value)
            if col == "Predicted_Attrition":
                if value == "Yes":
                    cell.fill = yes_fill
                elif value == "No":
                    cell.fill = no_fill
            if col in ("Age",):
                cell.alignment = center

    widths = {
        "EmployeeNumber": 14,
        "Age": 8,
        "Department": 22,
        "JobRole": 22,
        "MonthlyIncome": 12,
        "OverTime": 10,
        "Attrition": 10,
        "Predicted_Attrition": 18,
        "Predicted_Prob_Leave": 16,
        "Prediction_Match": 14,
    }
    for cidx, col in enumerate(cols, start=1):
        sheet.column_dimensions[get_column_letter(cidx)].width = widths.get(col, 14)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(df)+1}"


def main():
    if not DATA_PATH.exists():
        raise FileNotFoundError(f"Dataset not found: {DATA_PATH}")
    if not MODEL_PATH.exists():
        raise FileNotFoundError(f"Model not found: {MODEL_PATH}")

    df = load_dataset()
    pipe = joblib.load(MODEL_PATH)
    df_out, y_true, y_pred, y_prob = build_model_results(pipe, df)

    wb = Workbook()

    # ---- Predictions sheet (first so it is the active one) ----
    df_out_sorted = df_out.sort_values("Predicted_Prob_Leave", ascending=False)
    df_out_sorted = df_out_sorted.reset_index(drop=True)

    predictions_sheet = wb.active
    predictions_sheet.title = "Predictions"
    predictions_sheet.append(list(df_out_sorted.columns))
    for _, row in df_out_sorted.iterrows():
        predictions_sheet.append(list(row.values))

    style_workbook(wb, df_out_sorted)

    # ---- Summary sheet ----
    summary_sheet = wb.create_sheet("Prediction Summary")
    summary_sheet.append(["Metric", "Value"])
    summary_rows = [
        ("Total Employees", int(len(df_out))),
        ("Actual Leavers (Attrition = Yes)", int((df_out["Attrition"] == "Yes").sum())),
        ("Predicted Leavers", int((y_pred == 1).sum())),
        ("Prediction Match Rate", float((y_pred == y_true).mean())),
        ("Mean Predicted Probability of Leaving", float(np.mean(y_prob))),
        ("Max Predicted Probability of Leaving", float(np.max(y_prob))),
        ("Min Predicted Probability of Leaving", float(np.min(y_prob))),
        ("Mean Predicted Prob for Actual Leavers", float(np.mean(y_prob[y_true == 1]))),
        ("Mean Predicted Prob for Actual Non-Leavers", float(np.mean(y_prob[y_true == 0]))),
    ]
    for metric, value in summary_rows:
        summary_sheet.append([metric, value])

    for cell in summary_sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F3B63")
        cell.font = Font(bold=True, color="FFFFFF")

    summary_sheet.column_dimensions["A"].width = 46
    summary_sheet.column_dimensions["B"].width = 20

    path = EXPORT_PATH
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))

    print(f"Excel export saved to: {path}")
    print(f"Rows exported: {len(df_out)}")
    print(f"Summary metric rows: {len(summary_rows)}")


if __name__ == "__main__":
    main()
